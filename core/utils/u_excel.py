from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple, Union

from openpyxl import Workbook, load_workbook


ExcelPath = Union[str, Path]
ExcelRecord = Dict[Any, Any]


class ExcelReader:
    """读取 ``.xlsx`` 工作表，并转换为字典记录。"""

    def __init__(self, excel_path: ExcelPath, sheet_name: str = "Sheet1") -> None:
        self.file_path = Path(excel_path)
        self.sheet_name = sheet_name
        self._validate_file()
        self._headers, rows = self._read_sheet()
        self._records = self._build_records(self._headers, rows)

    @property
    def row_count(self) -> int:
        """返回工作表总行数，包含表头。"""
        return len(self._records) + (1 if self._headers else 0)

    @property
    def column_count(self) -> int:
        """返回工作表列数。"""
        return len(self._headers)

    def _validate_file(self) -> None:
        """校验文件是否存在以及扩展名是否受支持。"""
        if not self.file_path.is_file():
            raise FileNotFoundError(f"Excel文件不存在：{self.file_path}")
        if self.file_path.suffix.lower() != ".xlsx":
            raise ValueError(f"不支持的Excel文件格式：{self.file_path.suffix}")

    def _read_sheet(self) -> Tuple[List[Any], List[Sequence[Any]]]:
        """使用 openpyxl 读取 xlsx 工作表。"""
        workbook = load_workbook(
            self.file_path,
            read_only=True,
            data_only=True,
        )
        try:
            if self.sheet_name not in workbook.sheetnames:
                raise KeyError(f"工作表不存在：{self.sheet_name}")
            worksheet = workbook[self.sheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        return self._split_header_and_rows(rows)

    @classmethod
    def _split_header_and_rows(
        cls, rows: Sequence[Sequence[Any]]
    ) -> Tuple[List[Any], List[Sequence[Any]]]:
        """拆分并标准化工作表表头与数据行。"""
        if not rows:
            return [], []
        headers = [cls._normalize_cell(value) for value in rows[0]]
        return headers, list(rows[1:])

    @classmethod
    def _build_records(
        cls,
        headers: Sequence[Any],
        rows: Sequence[Sequence[Any]],
    ) -> List[ExcelRecord]:
        """按表头将每一行数据转换为字典。"""
        records = []
        for row in rows:
            normalized_row = [cls._normalize_cell(value) for value in row]
            records.append(dict(zip(headers, normalized_row)))
        return records

    @staticmethod
    def _normalize_cell(value: Any) -> Any:
        """保持旧数据驱动语义：空单元格为空串，整数数值转换为字符串。"""
        if value is None:
            return ""
        if isinstance(value, bool):
            return value
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return value

    def read_records(self) -> List[ExcelRecord]:
        """返回工作表中的全部数据记录，不包含表头。"""
        return [record.copy() for record in self._records]

    def get_row(self, row_number: int) -> Optional[ExcelRecord]:
        """按 Excel 行号读取一行；表头及之前的行返回 ``None``。"""
        if row_number <= 1:
            return None
        return self._records[row_number - 2].copy()

    def get_column(self, column_name: Any) -> List[Any]:
        """按表头名称读取整列数据。"""
        return [record[column_name] for record in self._records]

    def get_cell(self, row_number: int, column_name: Any) -> Any:
        """按 Excel 行号和表头名称读取单元格。"""
        row = self.get_row(row_number)
        return None if row is None else row[column_name]


class ExcelWriter:
    """将字典记录写入 ``.xlsx`` 工作表。"""

    def __init__(self, sheet_name: str = "Sheet1") -> None:
        self.sheet_name = sheet_name

    def write_records(
        self,
        records: List[Mapping[Any, Any]],
        excel_path: ExcelPath,
    ) -> None:
        """按照首条记录的字段顺序写入表头和数据。"""
        self._validate_records(records)
        output_path = Path(excel_path)
        if output_path.suffix.lower() != ".xlsx":
            raise ValueError(f"不支持的Excel文件格式：{output_path.suffix}")
        self._write_xlsx(records, output_path)

    @staticmethod
    def _validate_records(records: List[Mapping[Any, Any]]) -> None:
        """校验待写入数据必须是非空字典列表。"""
        if not isinstance(records, list):
            raise TypeError("数据必须是列表")
        if not records:
            raise ValueError("写入数据不能为空")
        if not all(isinstance(record, Mapping) for record in records):
            raise TypeError("列表中的每一项都必须是字典")

    def _write_xlsx(
        self,
        records: List[Mapping[Any, Any]],
        output_path: Path,
    ) -> None:
        """使用 openpyxl 写入 xlsx 文件。"""
        headers = list(records[0].keys())
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.sheet_name
        worksheet.append(headers)
        for record in records:
            worksheet.append([record.get(header) for header in headers])
        workbook.save(output_path)
        workbook.close()
